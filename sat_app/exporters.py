from __future__ import annotations

from io import BytesIO

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from sat_app.config import RISK_COLORS


def dataframe_to_excel_bytes(df: pd.DataFrame, sheet_name: str = "predicciones") -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
    buffer.seek(0)
    wb = load_workbook(buffer)
    ws = wb[sheet_name]
    header_fill = PatternFill("solid", fgColor="003D7A")
    risk_fills = {
        "ALTO": PatternFill("solid", fgColor="FEE2E2"),
        "MEDIO": PatternFill("solid", fgColor="FEF3C7"),
        "BAJO": PatternFill("solid", fgColor="DCFCE7"),
    }
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    risk_idx = None
    for idx, cell in enumerate(ws[1], start=1):
        if str(cell.value).lower() == "nivel_riesgo":
            risk_idx = idx
            break
    if risk_idx:
        for row in range(2, ws.max_row + 1):
            value = ws.cell(row=row, column=risk_idx).value
            fill = risk_fills.get(str(value))
            if fill:
                ws.cell(row=row, column=risk_idx).fill = fill
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def summary_pdf_bytes(df: pd.DataFrame, title: str, subtitle: str) -> bytes:
    buffer = BytesIO()
    with PdfPages(buffer) as pdf:
        fig, ax = plt.subplots(figsize=(8.27, 11.69))
        ax.axis("off")
        ax.text(0.02, 0.96, title, fontsize=20, fontweight="bold")
        ax.text(0.02, 0.92, subtitle, fontsize=11)

        total = len(df)
        counts = df["nivel_riesgo"].value_counts()
        lines = [
            f"Total estudiantes: {total}",
            f"Alto riesgo: {counts.get('ALTO', 0)}",
            f"Riesgo medio: {counts.get('MEDIO', 0)}",
            f"Bajo riesgo: {counts.get('BAJO', 0)}",
            f"Probabilidad promedio: {df['probabilidad'].mean():.2f}" if total else "Probabilidad promedio: 0.00",
        ]
        ax.text(0.02, 0.80, "\n".join(lines), fontsize=12, va="top")

        sample_cols = [col for col in ["id", "name", "student_key", "probabilidad", "nivel_riesgo"] if col in df.columns]
        sample = df[sample_cols].head(12)
        table_text = sample.to_string(index=False)
        ax.text(0.02, 0.64, "Top registros visibles", fontsize=13, fontweight="bold")
        ax.text(0.02, 0.60, table_text, fontsize=9, family="monospace", va="top")
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        fig2, ax2 = plt.subplots(figsize=(8.27, 5))
        ordered = ["ALTO", "MEDIO", "BAJO"]
        values = [counts.get(level, 0) for level in ordered]
        colors = [RISK_COLORS[level] for level in ordered]
        ax2.bar(ordered, values, color=colors)
        ax2.set_title("Distribucion de riesgo")
        ax2.set_ylabel("Estudiantes")
        pdf.savefig(fig2, bbox_inches="tight")
        plt.close(fig2)
    return buffer.getvalue()
